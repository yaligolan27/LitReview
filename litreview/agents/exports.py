"""Stage 17 — additional outputs: DOCX / Marp slides / NotebookLM podcast
(spec §9.17). All three are best-effort, each wrapped separately."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from ..core.context import RunContext
from ..core.marker_render import render_content_to_docx, set_rtl, to_plain_text
from ..core.state import SurveyState

_BADGE = {"HIGH": "🟢", "MODERATE": "🔵", "LIMITED": "🟠", "EMERGING": "🔴"}


def _out_dir(ctx: RunContext) -> Path:
    out = ctx.workdir / "outputs"
    out.mkdir(parents=True, exist_ok=True)
    return out


def build_docx(ctx: RunContext, state: SurveyState) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    doc = Document()
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)   # A4
        section.top_margin = section.bottom_margin = Cm(2.5)
        section.left_margin = section.right_margin = Cm(2.5)

    title = doc.add_heading(state.brief.topic, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(
        f"סקר ספרות · {date.today():%d.%m.%Y}"
        + (f" · {state.brief.author}" if state.brief.author else ""))
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if state.scorecard:
        doc.add_heading("Scorecard", level=1)
        p = doc.add_paragraph(
            f"ציון איכות: {state.scorecard.get('score', '—')}/100"
            + (" ⚠️ מתחת לסף" if state.scorecard.get("below_threshold") else ""))
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_rtl(p)

    if state.executive_summary:
        doc.add_heading("תקציר מנהלים", level=1)
        p = doc.add_paragraph(to_plain_text(state.executive_summary))
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_rtl(p)

    for i, sec in enumerate(state.sections, start=1):
        doc.add_heading(f"{i}. {sec.title}", level=1)
        badge = doc.add_paragraph(f"רמת אמינות: {sec.effective_confidence()}")
        badge.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_rtl(badge)
        render_content_to_docx(doc, sec.content)

    doc.add_heading("ביבליוגרפיה", level=1)
    for n, paper in enumerate(state.cited_papers, start=1):
        p = doc.add_paragraph(f"[{n}] {paper.apa}")
        p.paragraph_format.space_after = Pt(4)

    path = _out_dir(ctx) / f"survey_{date.today():%Y%m%d}.docx"
    doc.save(str(path))
    return str(path)


def build_slides(ctx: RunContext, state: SurveyState) -> str:
    lines = [
        "---", "marp: true", "theme: default", "paginate: true",
        'style: "section { direction: rtl; font-family: David, Arial; text-align: right; }"',
        "---", "",
        f"# {state.brief.topic}",
        f"### סקר ספרות · {len(state.cited_papers)} מקורות · "
        f"{len(state.sections)} פרקים", "",
    ]
    if state.scorecard:
        lines += [f"**ציון איכות: {state.scorecard.get('score', '—')}/100**", ""]
    for kpi in state.kpi_data[:4]:
        lines.append(f"- **{kpi['num']}** — {kpi['label']}")
    for i, sec in enumerate(state.sections, start=1):
        summary = to_plain_text(sec.content)[:250].rsplit(" ", 1)[0]
        lines += ["", "---", "",
                  f"## {i}. {sec.title} {_BADGE.get(sec.effective_confidence(), '')}",
                  "", summary + "…"]
    lines += ["", "---", "", "## מקורות",
              f"ביבליוגרפיה מלאה ({len(state.cited_papers)} מקורות) במסמך המלא."]
    path = _out_dir(ctx) / f"slides_{date.today():%Y%m%d}.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return str(path)


def build_podcast_text(ctx: RunContext, state: SurveyState) -> str:
    parts = [state.brief.topic, "=" * 30, ""]
    if state.executive_summary:
        parts += ["תקציר:", to_plain_text(state.executive_summary), ""]
    for i, sec in enumerate(state.sections, start=1):
        parts += [f"פרק {i}: {sec.title}", to_plain_text(sec.content), ""]
    parts += ["מקורות עיקריים:"]
    for paper in state.cited_papers[:15]:
        parts.append(f"- {paper.title} ({paper.year})")
    parts += ["", "הערת אמינות: הסקר נכתב בסיוע AI על בסיס מקורות אקדמיים "
                  "שאומתו ככל הניתן. פרטי האימות המלאים במסמך המקורי."]
    path = _out_dir(ctx) / f"notebooklm_{date.today():%Y%m%d}.txt"
    path.write_text("\n".join(parts), encoding="utf-8")
    return str(path)


def run_extras(ctx: RunContext, state: SurveyState) -> dict:
    results: dict[str, str] = {}
    try:
        results["docx"] = build_docx(ctx, state)
    except Exception as exc:  # noqa: BLE001 — best-effort per spec
        state.log("extras", f"DOCX failed: {exc}")
    if state.brief.output_slides:
        try:
            results["slides"] = build_slides(ctx, state)
        except Exception as exc:  # noqa: BLE001
            state.log("extras", f"slides failed: {exc}")
    if state.brief.output_podcast:
        try:
            results["podcast"] = build_podcast_text(ctx, state)
        except Exception as exc:  # noqa: BLE001
            state.log("extras", f"podcast failed: {exc}")
    state.log("extras", "additional outputs", **{k: str(v) for k, v in results.items()})
    return results
