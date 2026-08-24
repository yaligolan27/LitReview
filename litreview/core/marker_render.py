"""Markers → Word / plain text (spec §9.17, core/marker_render.py).

Word does not run MathJax, so LaTeX is converted to readable Unicode:
Greek letters, operators, \\frac→(a)/(b), \\sqrt→√(), powers and indices to
super/subscripts, unknown commands dropped.
"""

from __future__ import annotations

import re

GREEK = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε",
    "zeta": "ζ", "eta": "η", "theta": "θ", "iota": "ι", "kappa": "κ",
    "lambda": "λ", "mu": "μ", "nu": "ν", "xi": "ξ", "pi": "π", "rho": "ρ",
    "sigma": "σ", "tau": "τ", "upsilon": "υ", "phi": "φ", "chi": "χ",
    "psi": "ψ", "omega": "ω",
    "Gamma": "Γ", "Delta": "Δ", "Theta": "Θ", "Lambda": "Λ", "Xi": "Ξ",
    "Pi": "Π", "Sigma": "Σ", "Phi": "Φ", "Psi": "Ψ", "Omega": "Ω",
}

OPERATORS = {
    "leq": "≤", "geq": "≥", "neq": "≠", "approx": "≈", "sim": "∼",
    "times": "×", "cdot": "·", "div": "÷", "pm": "±", "mp": "∓",
    "infty": "∞", "sum": "∑", "prod": "∏", "int": "∫", "partial": "∂",
    "nabla": "∇", "propto": "∝", "in": "∈", "subset": "⊂", "cup": "∪",
    "cap": "∩", "rightarrow": "→", "leftarrow": "←", "Rightarrow": "⇒",
    "forall": "∀", "exists": "∃", "degree": "°", "circ": "°",
}

_SUPERSCRIPTS = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")
_SUBSCRIPTS = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")


def latex_to_unicode(latex: str) -> str:
    text = latex.strip()
    for _ in range(4):   # nested fractions, up to 4 levels
        new = re.sub(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", text)
        if new == text:
            break
        text = new
    text = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", text)
    for name, char in {**GREEK, **OPERATORS}.items():
        text = re.sub(rf"\\{name}\b", char, text)
    text = re.sub(r"\\(?:operatorname|text|mathrm|mathbf|mathit)\{([^{}]*)\}",
                  r"\1", text)

    def _sup(match: re.Match) -> str:
        # group(1) is the braced form ('' for empty braces), group(2) the \w
        # form; `group(1) or group(2)` collapsed ''->None and crashed on None.
        content = match.group(1) if match.group(1) is not None else (match.group(2) or "")
        if not content:
            return ""
        translated = content.translate(_SUPERSCRIPTS)
        return translated if translated != content or content.isdigit() \
            else f"^({content})"

    def _sub(match: re.Match) -> str:
        content = match.group(1) if match.group(1) is not None else (match.group(2) or "")
        if not content:
            return ""
        translated = content.translate(_SUBSCRIPTS)
        return translated if all(ch in "0123456789+-=()" for ch in content) \
            else f"_({content})"

    text = re.sub(r"\^\{([^{}]*)\}|\^(\w)", _sup, text)
    text = re.sub(r"_\{([^{}]*)\}|_(\w)", _sub, text)
    text = re.sub(r"\\[a-zA-Z]+", "", text)   # drop unknown commands
    text = text.replace("{", "").replace("}", "")
    return re.sub(r"\s+", " ", text).strip()


def to_plain_text(content: str) -> str:
    """Markers → clean prose (podcast / slides / previews)."""
    text = content
    text = re.sub(r"\[FORMULA\](.*?)\[/FORMULA\]",
                  lambda m: latex_to_unicode(m.group(1)), text, flags=re.DOTALL)
    text = re.sub(r"\\\((.*?)\\\)", lambda m: latex_to_unicode(m.group(1)), text)
    text = re.sub(r"\[CALLOUT(?::[a-zA-Z]+)?\](.*?)\[/CALLOUT\]", r"\1", text,
                  flags=re.DOTALL)
    text = re.sub(r"\[CASE\](.*?)\[/CASE\]",
                  lambda m: "מקרה בוחן: " + m.group(1).replace("|", " — "),
                  text, flags=re.DOTALL)
    text = re.sub(r"\[EXAMPLE\](.*?)\[/EXAMPLE\]",
                  lambda m: "דוגמה מחושבת: " + " — ".join(
                      p.strip() for p in m.group(1).split("|") if p.strip()),
                  text, flags=re.DOTALL)
    text = re.sub(r"\[KPI\](.*?)\[/KPI\]", "", text, flags=re.DOTALL)
    text = re.sub(r"\[TABLE\](.*?)\[/TABLE\]", "", text, flags=re.DOTALL)
    text = re.sub(r"\[TRL:(\d)\]", r"(TRL \1)", text)
    text = re.sub(r"\[(?:KPI_DATA|CONCLUSION|ROI_CALC)\]|\[/(?:KPI_DATA|CONCLUSION|ROI_CALC)\]",
                  "", text)
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = text.replace("**", "")
    text = text.replace("[BEST]", "").replace("[BAD]", "")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ---------------------------------------------------------------------------
# Word rendering (python-docx is an optional dependency)
# ---------------------------------------------------------------------------


def render_content_to_docx(doc, content: str) -> None:
    """Render chapter markdown+markers into a python-docx Document."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    def rtl(paragraph):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_rtl(paragraph)

    segments = re.split(
        r"(\[(?:FORMULA|TABLE|CALLOUT[^\]]*|CASE|KPI|EXAMPLE)\].*?"
        r"\[/(?:FORMULA|TABLE|CALLOUT|CASE|KPI|EXAMPLE)\])",
        content, flags=re.DOTALL)
    for segment in segments:
        segment = segment.strip()
        if not segment:
            continue
        if segment.startswith("[FORMULA]"):
            inner = re.sub(r"^\[FORMULA\]|\[/FORMULA\]$", "", segment, flags=re.DOTALL)
            p = doc.add_paragraph(latex_to_unicode(inner))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif segment.startswith("[TABLE]"):
            inner = re.sub(r"^\[TABLE\]|\[/TABLE\]$", "", segment, flags=re.DOTALL)
            _add_table(doc, inner)
        elif segment.startswith("[CALLOUT"):
            inner = re.sub(r"^\[CALLOUT(?::[a-zA-Z]+)?\]|\[/CALLOUT\]$", "",
                           segment, flags=re.DOTALL)
            p = doc.add_paragraph()
            run = p.add_run(to_plain_text(inner))
            run.bold = True
            rtl(p)
        elif segment.startswith(("[CASE]", "[KPI]", "[EXAMPLE]")):
            p = doc.add_paragraph(to_plain_text(segment))
            rtl(p)
        else:
            for line in segment.splitlines():
                line = line.strip()
                if not line:
                    continue
                heading = re.match(r"^(#{2,4})\s*(.*)$", line)
                if heading:
                    doc.add_heading(to_plain_text(heading.group(2)),
                                    level=min(4, len(heading.group(1))))
                    continue
                p = doc.add_paragraph(to_plain_text(line))
                rtl(p)


def _add_table(doc, markdown: str) -> None:
    rows = [r.strip() for r in markdown.strip().splitlines() if r.strip()]
    rows = [r for r in rows if not re.fullmatch(r"\|?[\s\-:|]+\|?", r)]
    if not rows:
        return
    matrix = [[c.strip().replace("[BEST]", "").replace("[BAD]", "").strip()
               for c in r.strip("|").split("|")] for r in rows]
    cols = max(len(r) for r in matrix)
    table = doc.add_table(rows=len(matrix), cols=cols)
    try:
        table.style = "Light Grid Accent 1"
    except KeyError:
        pass
    for i, row in enumerate(matrix):
        for j, cell in enumerate(row):
            table.rows[i].cells[j].text = cell


def set_rtl(paragraph) -> None:
    """Real Word RTL: inject <w:bidi/> into the paragraph properties."""
    from docx.oxml.ns import qn

    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        bidi = pPr.makeelement(qn("w:bidi"), {})
        pPr.append(bidi)
